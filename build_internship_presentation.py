from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

OUT = "assets/files/Binod_Raj_Regmi_Internship_Presentation.pptx"
NOTES = "assets/files/Binod_Raj_Regmi_Internship_Speaker_Notes.md"

NAVY = RGBColor(12, 31, 53)
NAVY2 = RGBColor(21, 48, 77)
ORANGE = RGBColor(224, 107, 45)
GREY = RGBColor(224, 229, 232)
MID = RGBColor(104, 119, 130)
INK = RGBColor(37, 48, 57)
WHITE = RGBColor(255, 255, 255)
PALE = RGBColor(245, 247, 248)

slides = [
    ("Internship Presentation", ["Infrastructure Department", "Tarakeshwor Municipality"], "Binod Raj Regmi | BE Civil | NCIT | 23 April–23 July 2026", None),
    ("What • How • Where", ["WHAT  Municipal context + internship purpose", "HOW  Infrastructure Department • office work • field visits", "WHERE  Dharmasthali, Kathmandu • Wards 6, 7 and 8 • project sites", "OUTPUT  Documents • observations • estimates • learning"], "ROADMAP", "roadmap"),
    ("Internship Overview", ["Three-month BE Civil internship", "Host: Tarakeshwor Municipality", "Department: Infrastructure Department", "Period: 23 April 2026 to 23 July 2026", "Focus: connect classroom knowledge with municipal practice"], "OVERVIEW", None),
    ("About Tarakeshwor Municipality", ["Local government institution in Kathmandu District, Bagmati Province", "Plans and delivers local public services and development", "Coordinates elected representatives, officials and technical personnel", "Infrastructure scope: roads, public buildings, drainage and services"], "MUNICIPAL CONTEXT", None),
    ("Infrastructure Department", ["Plans, estimates, supervises and monitors infrastructure works", "Supports drawings, quantities, rates, measurement and valuation", "Maintains technical records and project documentation", "Coordinates with wards, contractors, consultants and communities"], "ENGINEERING FUNCTION", None),
    ("Internship Objectives", ["Relate BE Civil theory to field conditions", "Understand municipal project procedures and records", "Build skill in estimation, rate analysis and drawing review", "Observe construction quality, safety and coordination", "Develop professional discipline and communication"], "PURPOSE", None),
    ("Activities Carried Out", ["Orientation and municipal documentation", "Field visits: Old Age Home, gabion wall, road marking", "Rate analysis using Excel, norms and district rates", "Truss-based document-storage structure estimate", "Technical learning through supervised practice"], "ACTIVITY MAP", None),
    ("Orientation & Municipal Documentation", ["Observed department procedures and project workflow", "Reviewed proposals, BOQs, estimates, valuations and drawings", "Reviewed records related to map renewal and property transfer", "Understood construction approval and document coordination", "Practised workplace discipline and confidentiality"], "ROLE: OBSERVED • REVIEWED", "photo"),
    ("Old Age Home | Ward 8", ["Assisted surveyors with tape measurement and site marking", "Used cadastral map information for boundary verification", "Recorded observations in the field book", "Observed RCC retaining compound-wall work", "Checked dimensions under technical supervision"], "ROLE: ASSISTED • CHECKED • OBSERVED", "photo"),
    ("Gabion Wall | Ward 6", ["Observed stepped foundation and gabion-wall construction", "Observed mesh-box assembly, lacing and stone filling", "Checked structure dimensions with the supervisor", "Understood roadside erosion and embankment protection", "Observed barriers around an active roadside work area"], "ROLE: OBSERVED • CHECKED", "photo"),
    ("Road Marking | Wards 7 & 8", ["Observed marking on newly constructed flexible roads", "Yellow edge lines and black-and-white zebra crossings", "Observed sweeping, string alignment and machine application", "Checked marking dimensions under technical guidance", "Noted the importance of dry, clean pavement"], "ROLE: OBSERVED • CHECKED", "photo"),
    ("Rate Analysis | Purpose & Workflow", ["Purpose: calculate a realistic unit cost", "Identify item, unit and resource requirements", "Apply DUDBC and Department of Roads norms", "Price with Kathmandu District approved rates 2081/82", "Use market rates only when an item is unavailable"], "WORKFLOW", "flow"),
    ("Rate Analysis | Tools & Contribution", ["Prepared analyses in Microsoft Excel and municipal template", "Entered quantities, rates and cost components", "Applied formulas, checked subtotals and organized sheets", "Referred to DUDBC and Department of Roads norms", "Submitted work for supervisor review and correction"], "ROLE: PREPARED • CHECKED • REVIEWED", "photo"),
    ("Truss Estimate | Project Context", ["Proposed document-storage hall above an existing masonry building", "Approximate floor area: 12.5 m × 5.5 m", "Existing mild-steel truss roof to be removed and reinstalled", "Estimate covered floor system, supports, walls, openings and stair", "Assignment: prepare and compare two initial-cost alternatives"], "ASSIGNMENT", "photo"),
    ("Truss Alternatives | Result", ["Alternative A: RCC slab floor with existing truss reused as roof", "Alternative B: 6 mm mild-steel plate floor with steel supports", "Both options included truss removal and reinstallation", "Prepared both alternatives in one Excel worksheet", "Initial-cost comparison: RCC slab was less costly"], "ROLE: MEASURED • PREPARED • COMPARED", "compare"),
    ("Technical Tools & Methods", ["Microsoft Excel: quantities, formulas and cost comparison", "AutoCAD: drawing study, interpretation and scaling", "Google Maps + KML: location and coordinate information", "METWET: steel-section and plate unit weights", "Field book, tape measurement and supervised site checks"], "TOOLKIT", None),
    ("Practical & Professional Skills", ["Site-boundary verification and map interpretation", "Construction observation and dimensional checking", "Estimation, rate analysis and technical documentation", "Communication with technical staff and landowners", "Teamwork, time management and workplace discipline"], "LEARNING", None),
    ("Challenges & Lessons Learned", ["Finding matching items across norms and rate schedules", "Checking units, specifications, quantities and rates", "Working around dust, wet pavement and active traffic", "Field decisions must respond to actual site conditions", "Accurate records and supervisor review protect quality"], "REFLECTION", None),
    ("Conclusion & Recommendations", ["The internship connected BE Civil study with municipal practice", "It strengthened estimation, field observation and documentation skills", "Future work should maintain careful checking and records", "Continue learning from site supervision and technical review", "Use approved norms and rates with clear basis and judgement"], "TAKEAWAY", None),
    ("Thank You", ["Questions?"], "Binod Raj Regmi | BE Civil | NCIT", "thankyou"),
]

notes = [
    "Introduce yourself, NCIT and the host department. State the internship period and explain that the presentation focuses on what you encountered, how you participated, and what you learned.",
    "Use this as the presentation map: the municipal setting is the context, the Infrastructure Department and office/field activities are the method, and Tarakeshwor plus the project wards are the locations. The outputs were reviewed documents, field observations and estimates.",
    "This was a three-month internship completed as part of BE Civil at NCIT. Emphasize exposure to both office-based and field-based municipal engineering activities.",
    "Tarakeshwor Municipality is a local government institution in Kathmandu District. Its infrastructure responsibilities affect transport, public services, safety and local development.",
    "Explain the department as the technical link between approved municipal plans and physical infrastructure. Mention planning, estimation, supervision, records and coordination without implying independent authority.",
    "The objectives were practical: connect theory to real work, understand municipal procedures, improve estimation and drawing skills, and learn professional conduct under supervision.",
    "Give the audience the sequence: orientation and records first, then field visits, then rate analysis and the truss estimate. The level of involvement varied by activity.",
    "During the first two weeks, your role was mainly to observe and review existing records. Mention proposals, BOQs, estimates, valuations and drawings, plus map renewal, property transfer and approval procedures.",
    "At the Ward 8 Old Age Home site, you assisted with tape measurement, marking and boundary verification using cadastral information. You observed RCC retaining compound-wall reinforcement and formwork, and checked dimensions under supervision; you did not independently approve the work.",
    "At Ward 6, the gabion wall addressed roadside erosion and embankment stability. You observed stepped construction, mesh boxes, lacing and stone filling, and checked dimensions with the supervisor. Keep the focus on observation and site-specific protection.",
    "At Wards 7 and 8, you observed thermoplastic edge lines and zebra crossings on flexible roads. Surface preparation, string alignment, machine application and dry conditions mattered. Your role was mainly observation and dimension checking under guidance.",
    "Explain rate analysis as the unit-cost building block of an estimate. The workflow moved from identifying the item and unit to selecting norms, applying approved rates and summing resource costs. Market rates were used only for unavailable items.",
    "You personally entered data, applied formulas, organized cost breakdowns and corrected work after review. Name the sources precisely: DUDBC norms, Department of Roads norms and Kathmandu District approved rates for 2081/82.",
    "The assignment was a document-storage hall above an existing one-storey stone masonry building. Explain that the estimate had to account for the existing truss and the proposed upper structure, not just a floor surface.",
    "Compare the two options at a high level. RCC used the slab alternative; the other used a 6 mm mild-steel plate and steel support system. Both reused the existing truss as roof. The estimate you prepared indicated RCC was less costly initially. Do not state totals that are not on the slide.",
    "This slide gathers the working toolkit. Excel supported calculations; AutoCAD supported drawing study and scaling; Google Maps and KML supported location data; METWET supported steel weights; field tools supported supervised checks.",
    "Separate technical skills from professional skills. The experience improved observation, documentation, communication, coordination and discipline, alongside estimation and rate-analysis ability.",
    "The main challenge in rate analysis was locating matching items and verifying units and specifications. Field work also depended on dust, wet surfaces, traffic and site conditions. The lesson was to check carefully and seek technical review.",
    "Conclude that the internship bridged classroom knowledge and municipal engineering practice. Recommend continued attention to approved references, accurate records, field supervision and learning through review.",
    "Thank the audience and invite questions. This is also the natural point to mention that site photographs and personal-work photographs can be added to the marked slides.",
]

assert len(slides) == 20 and len(notes) == 20

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def add_text(slide, text, x, y, w, h, size=18, color=INK, bold=False, font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def rect(slide, x, y, w, h, fill, line=None, radius=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def line(slide, x1, y1, x2, y2, color=GREY, width=1.0):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    return shape


def base(slide, section, number):
    rect(slide, 0, 0, 13.333, 7.5, WHITE)
    rect(slide, 0, 0, 0.18, 7.5, ORANGE)
    rect(slide, 0.18, 0, 13.153, 0.12, NAVY)
    for i in range(12):
        line(slide, 9.5 + i * 0.33, 0.25, 9.5 + i * 0.33, 1.35, GREY, 0.45)
    for i in range(5):
        line(slide, 9.35, 0.35 + i * 0.22, 13.0, 0.35 + i * 0.22, GREY, 0.45)
    add_text(slide, section, 0.65, 0.42, 4.5, 0.25, 9, ORANGE, True)
    add_text(slide, f"{number:02d}", 12.25, 0.38, 0.55, 0.28, 10, NAVY, True, align=PP_ALIGN.RIGHT)


def bullets(slide, items, x=0.85, y=1.55, w=7.6, h=4.9, size=20):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(0.04); tf.margin_right = Inches(0.04)
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = "- " + item
        p.level = 0
        p.space_after = Pt(13)
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = INK
    return box


def tag(slide, text, x=0.85, y=6.65, w=3.25):
    rect(slide, x, y, w, 0.36, PALE, GREY, True)
    add_text(slide, text, x + 0.13, y + 0.08, w - 0.25, 0.18, 8.5, MID, True)


def placeholder(slide, label="[Insert site photograph]"):
    rect(slide, 9.0, 1.55, 3.35, 3.65, PALE, MID, True)
    line(slide, 9.0, 1.55, 12.35, 5.2, GREY, 0.8)
    line(slide, 12.35, 1.55, 9.0, 5.2, GREY, 0.8)
    add_text(slide, label, 9.35, 3.1, 2.65, 0.7, 16, NAVY, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def roadmap(slide):
    labels = [("WHAT", "Municipal context\n+ purpose"), ("HOW", "Department\n+ activities"), ("WHERE", "Dharmasthali\n+ project sites")]
    xs = [0.85, 4.55, 8.25]
    for i, (head, body) in enumerate(labels):
        rect(slide, xs[i], 2.15, 3.0, 1.9, NAVY if i == 1 else PALE, NAVY, True)
        add_text(slide, head, xs[i] + 0.22, 2.43, 2.55, 0.28, 12, ORANGE if i == 1 else ORANGE, True, align=PP_ALIGN.CENTER)
        add_text(slide, body, xs[i] + 0.25, 2.86, 2.5, 0.7, 18, WHITE if i == 1 else INK, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        if i < 2:
            line(slide, xs[i] + 3.0, 3.1, xs[i] + 3.7, 3.1, ORANGE, 2.2)
    rect(slide, 2.25, 5.05, 8.85, 0.75, ORANGE, ORANGE, True)
    add_text(slide, "OUTPUTS  documents  •  observations  •  estimates  •  learning", 2.55, 5.28, 8.25, 0.22, 13, WHITE, True, align=PP_ALIGN.CENTER)


def flow(slide):
    labels = ["Item + unit", "Norms", "Approved rates", "Cost breakdown", "Review"]
    x = 8.8
    for i, item in enumerate(labels):
        y = 1.65 + i * 0.72
        rect(slide, x, y, 2.75, 0.48, NAVY if i == 4 else PALE, NAVY, True)
        add_text(slide, item, x + 0.1, y + 0.13, 2.55, 0.2, 12, WHITE if i == 4 else INK, True, align=PP_ALIGN.CENTER)
        if i < 4: line(slide, x + 1.38, y + 0.48, x + 1.38, y + 0.72, ORANGE, 1.5)


def comparison(slide):
    for x, head, body, active in [(8.55, "RCC SLAB", "Floor: RCC\nRoof: reused truss", True), (10.35, "MS PLATE", "Floor: 6 mm plate\nRoof: reused truss", False)]:
        rect(slide, x, 1.75, 1.55, 2.25, NAVY if active else PALE, NAVY, True)
        add_text(slide, head, x + 0.08, 2.0, 1.39, 0.35, 10, ORANGE, True, align=PP_ALIGN.CENTER)
        add_text(slide, body, x + 0.12, 2.65, 1.31, 0.7, 12, WHITE if active else INK, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    line(slide, 10.1, 2.87, 10.3, 2.87, ORANGE, 2)
    add_text(slide, "LESS COSTLY\n(initial comparison)", 8.8, 4.45, 2.9, 0.58, 14, ORANGE, True, align=PP_ALIGN.CENTER)

for number, (title, items, section, visual) in enumerate(slides, start=1):
    slide = prs.slides.add_slide(blank)
    if number == 1:
        rect(slide, 0, 0, 13.333, 7.5, NAVY)
        rect(slide, 0, 0, 0.25, 7.5, ORANGE)
        for i in range(11):
            line(slide, 8.0 + i * 0.42, 0.6, 8.0 + i * 0.42, 6.9, NAVY2, 0.65)
        for i in range(12):
            line(slide, 7.8, 0.65 + i * 0.48, 12.9, 0.65 + i * 0.48, NAVY2, 0.65)
        add_text(slide, "INTERNSHIP REPORT DEFENSE", 0.9, 1.0, 5.8, 0.3, 12, ORANGE, True)
        add_text(slide, title, 0.85, 1.75, 7.1, 1.1, 34, WHITE, True, font="Aptos Display")
        add_text(slide, "Infrastructure Department\nTarakeshwor Municipality", 0.9, 3.15, 5.7, 0.95, 22, GREY, False)
        add_text(slide, items[0], 0.9, 5.35, 5.7, 0.3, 15, WHITE, True)
        add_text(slide, items[1], 0.9, 5.78, 6.2, 0.3, 15, WHITE, False)
        add_text(slide, section, 0.9, 6.7, 5.6, 0.2, 10, ORANGE, True)
    else:
        base(slide, section, number)
        add_text(slide, title, 0.82, 0.85, 8.2, 0.6, 28, NAVY, True, font="Aptos Display")
        if number == 2:
            roadmap(slide)
        elif number == 20:
            rect(slide, 0.85, 2.0, 7.0, 2.2, NAVY, NAVY, True)
            add_text(slide, "Questions?", 1.2, 2.55, 6.3, 0.75, 34, WHITE, True, font="Aptos Display", align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
            add_text(slide, items[0], 0.9, 5.55, 6.5, 0.35, 16, INK, True)
            add_text(slide, section, 0.9, 6.65, 5.5, 0.2, 10, ORANGE, True)
        else:
            bullets(slide, items, size=18 if len(items) >= 5 else 20)
            if visual == "photo": placeholder(slide)
            elif visual == "roadmap": roadmap(slide)
            elif visual == "flow": flow(slide)
            elif visual == "compare": comparison(slide)
            if section:
                tag(slide, section)

prs.save(OUT)

notes_lines = ["# Internship Presentation Speaker Notes", "", "Student: Binod Raj Regmi  ", "Program: BE Civil, NCIT  ", "Host: Tarakeshwor Municipality, Infrastructure Department  ", "", "These notes are separate from the PowerPoint and are editable. Keep the delivery close to 15 minutes.", ""]
for i, ((title, _, _, _), note) in enumerate(zip(slides, notes), start=1):
    notes_lines.extend([f"## Slide {i}: {title}", note, ""])
notes_lines.extend(["## Photo and Edit Checklist", "- Slides 8–11: add relevant municipal office or site photographs.", "- Slides 13–15: add personal work screenshots or photographs of Excel/estimate sheets where permitted.", "- Replace only the marked editable placeholders; do not add stock or generated images.", "- Confirm any added image captions and project details against the report."])
with open(NOTES, "w", encoding="utf-8") as handle:
    handle.write("\n".join(notes_lines))

print(OUT)
print(NOTES)
print(len(prs.slides))
